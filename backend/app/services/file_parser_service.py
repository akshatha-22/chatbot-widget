import csv
import logging
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Callable, Dict, List, Optional

logger = logging.getLogger(__name__)

PAGE_MARKER_RE = re.compile(r"\[PAGE\s+\d+\]", re.IGNORECASE)

MAX_PDF_OCR_BYTES = 100 * 1024 * 1024
PDF_FILE_API_THRESHOLD_BYTES = 4 * 1024 * 1024
MIN_TEXT_LENGTH = 20
# Process every page that still needs OCR after local extractors (flipbooks often need 100+).
MAX_OCR_PAGES = 512
MAX_CHARS_PER_PAGE = 3000
PYMUPDF_MAX_WORKERS = 4
OCR_MAX_WORKERS = 4

ProgressCallback = Callable[[str], None]


def is_row_chunked_format(filename: str) -> bool:
    """True for spreadsheet formats indexed one row per chunk."""
    ext = os.path.splitext(filename.lower())[1]
    return ext in (".xlsx", ".xls", ".csv")


def _format_row_chunk(headers: List[str], values: tuple) -> Optional[str]:
    """Format one spreadsheet row as a labeled chunk string."""
    parts: List[str] = []
    for idx, value in enumerate(values):
        if value is None or str(value).strip() == "":
            cell = "N/A"
        else:
            cell = str(value).strip()
        label = headers[idx] if idx < len(headers) and headers[idx] else f"Column {idx + 1}"
        parts.append(f"{label}: {cell}")
    if not parts:
        return None
    if all(
        (value is None or str(value).strip() == "")
        for value in values
    ):
        return None
    return " | ".join(parts)


def parse_row_chunks(file_path: str, filename: str) -> Optional[List[str]]:
    """
    Parse XLSX/XLS/CSV row-by-row. Each non-empty row becomes exactly one chunk.
    Returns None for non-tabular formats (caller uses character-based chunking).
    """
    if not is_row_chunked_format(filename):
        return None

    ext = os.path.splitext(filename.lower())[1]

    try:
        if ext == ".csv":
            chunks: List[str] = []
            with open(file_path, "r", encoding="utf-8", errors="ignore", newline="") as handle:
                reader = csv.reader(handle)
                rows = list(reader)
            if not rows:
                return []

            headers = [h.strip() or f"Column {i + 1}" for i, h in enumerate(rows[0])]
            data_rows = rows[1:] if len(rows) > 1 else rows
            if not any(any(cell.strip() for cell in row) for row in data_rows):
                return []

            for row in data_rows:
                padded = tuple(row) + tuple("" for _ in range(max(0, len(headers) - len(row))))
                chunk = _format_row_chunk(headers, padded[: len(headers)])
                if chunk:
                    chunks.append(chunk)
            return chunks

        if ext in (".xlsx", ".xls"):
            import openpyxl

            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            chunks = []
            for sheet in wb.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                first = rows[0]
                has_header = all(isinstance(cell, str) for cell in first if cell is not None)
                if has_header and len(rows) > 1:
                    headers = [
                        str(cell).strip() if cell is not None else f"Column {i + 1}"
                        for i, cell in enumerate(first)
                    ]
                    data_rows = rows[1:]
                else:
                    width = max(len(first), 1)
                    headers = [f"Column {i + 1}" for i in range(width)]
                    data_rows = rows

                for row in data_rows:
                    padded = tuple(row) + tuple(None for _ in range(max(0, len(headers) - len(row))))
                    chunk = _format_row_chunk(headers, padded[: len(headers)])
                    if chunk:
                        chunks.append(chunk)
            return chunks
    except Exception as exc:
        raise ValueError(f"Error parsing structured file '{filename}': {exc}") from exc

    return None


def _ensure_pdf_page_markers(text: str) -> str:
    """Wrap OCR or legacy extractor output with [PAGE N] when markers are missing."""
    cleaned = (text or "").strip()
    if not cleaned:
        return cleaned
    if PAGE_MARKER_RE.search(cleaned):
        return cleaned
    return f"[PAGE 1]\n{cleaned}"


def _get_pdf_page_count(file_path: str) -> int:
    try:
        import pdfplumber

        with pdfplumber.open(file_path) as pdf:
            return len(pdf.pages)
    except Exception:
        try:
            import fitz

            with fitz.open(file_path) as doc:
                return len(doc)
        except Exception:
            try:
                import PyPDF2

                with open(file_path, "rb") as handle:
                    return len(PyPDF2.PdfReader(handle).pages)
            except Exception:
                return 0


def _format_tables(tables: list) -> str:
    """Convert pdfplumber table arrays to readable pipe-delimited text."""
    output: List[str] = []
    for table in tables:
        for row in table:
            cleaned = [str(cell).strip() if cell else "" for cell in row]
            if any(cleaned):
                output.append(" | ".join(cleaned))
    return "\n".join(output)


def _missing_pages(all_pages: Dict[int, str], total_pages: int) -> List[int]:
    """Pages with no extractable text at all (for local extractors)."""
    return [
        page_num
        for page_num in range(1, total_pages + 1)
        if page_num not in all_pages or not (all_pages.get(page_num) or "").strip()
    ]


def _ocr_candidate_pages(all_pages: Dict[int, str], total_pages: int) -> List[int]:
    """Pages that need Gemini OCR — truly empty or below minimum text threshold."""
    return [
        page_num
        for page_num in range(1, total_pages + 1)
        if page_num not in all_pages
        or len((all_pages.get(page_num) or "").strip()) < MIN_TEXT_LENGTH
    ]


def _truncate_dense_pages(all_pages: Dict[int, str]) -> None:
    """Cap per-page text before chunking to keep embedding batches efficient."""
    for page_num, content in list(all_pages.items()):
        if len(content) > MAX_CHARS_PER_PAGE:
            all_pages[page_num] = content[:MAX_CHARS_PER_PAGE]
            logger.info(
                "Page %s truncated to %s chars in %s",
                page_num,
                MAX_CHARS_PER_PAGE,
                "pdf",
            )


def _ocr_page_limit(total_pages: int, ocr_needed: int) -> int:
    """How many pages to OCR — all candidates, bounded by doc size and safety cap."""
    if ocr_needed <= 0:
        return 0
    return min(ocr_needed, total_pages, MAX_OCR_PAGES)


def _gemini_ocr_pages_parallel(
    file_path: str,
    page_nums: List[int],
    on_progress: Optional[ProgressCallback] = None,
) -> Dict[int, str]:
    """OCR multiple scanned pages in parallel with progress updates."""
    if not page_nums:
        return {}

    results: Dict[int, str] = {}
    total = len(page_nums)
    completed = 0

    with ThreadPoolExecutor(max_workers=OCR_MAX_WORKERS) as executor:
        futures = {
            executor.submit(_gemini_ocr_page, file_path, page_num): page_num
            for page_num in page_nums
        }
        for future in as_completed(futures):
            page_num = futures[future]
            try:
                text = future.result()
            except Exception as exc:
                logger.warning("Gemini OCR page %s failed: %s", page_num, exc)
                text = ""
            if text:
                results[page_num] = text
            completed += 1
            if on_progress and (completed == 1 or completed == total or completed % 10 == 0):
                on_progress(f"OCR page {completed} of {total}…")

    return results


def _gemini_ocr_page(file_path: str, page_num: int) -> str:
    """Use Gemini vision to OCR a single scanned page (last resort)."""
    from app.config import settings

    if not settings.gemini_configured():
        return ""

    try:
        import fitz
        from google import genai
        from google.genai import types

        with fitz.open(file_path) as doc:
            page = doc[page_num - 1]
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("png")

        client = genai.Client(api_key=settings.GEMINI_API_KEY.strip())
        model = settings.GEMINI_MODEL.removeprefix("models/")
        response = client.models.generate_content(
            model=model,
            contents=[
                types.Content(
                    role="user",
                    parts=[
                        types.Part.from_bytes(data=img_bytes, mime_type="image/png"),
                        types.Part.from_text(
                            text=(
                                "Extract all text from this page exactly as it appears. "
                                "Return only the text, no commentary."
                            )
                        ),
                    ],
                )
            ],
        )
        return (getattr(response, "text", None) or "").strip()
    except Exception as exc:
        logger.warning("Gemini OCR page %s failed: %s", page_num, exc)
        return ""


def _extract_pymupdf_pages_parallel(
    file_path: str, page_nums: List[int]
) -> Dict[int, str]:
    """Extract missing pages via PyMuPDF using a small thread pool."""
    if not page_nums:
        return {}

    try:
        import fitz
    except Exception as exc:
        logger.warning("PyMuPDF unavailable: %s", exc)
        return {}

    results: Dict[int, str] = {}
    try:
        doc = fitz.open(file_path)

        def extract_page_fitz(page_num: int) -> tuple[int, str]:
            try:
                text = doc[page_num - 1].get_text("text").strip()
                return page_num, text
            except Exception as exc:
                logger.warning("PyMuPDF page %s failed: %s", page_num, exc)
                return page_num, ""

        with ThreadPoolExecutor(max_workers=PYMUPDF_MAX_WORKERS) as executor:
            futures = {
                executor.submit(extract_page_fitz, page_num): page_num
                for page_num in page_nums
            }
            for future in as_completed(futures):
                page_num, text = future.result()
                if text:
                    results[page_num] = text

        doc.close()
    except Exception as exc:
        logger.warning("PyMuPDF failed: %s", exc)

    return results


def _extract_pdf_deep(
    file_path: str,
    filename: str,
    on_progress: Optional[ProgressCallback] = None,
) -> str:
    """
    Deep extraction — reads every page with up to four strategies per missing page.
    Gemini OCR is capped and only used for genuinely empty/scanned pages.
    """
    total_pages = _get_pdf_page_count(file_path)
    if total_pages == 0:
        raise ValueError(f"Could not determine page count for '{filename}'")

    if on_progress:
        on_progress(f"Reading {total_pages} pages…")

    all_pages: Dict[int, str] = {}

    # Strategy 1 — pdfplumber (text + tables, all pages in one pass)
    try:
        import pdfplumber

        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                try:
                    page_text = page.extract_text() or ""
                    table_text = _format_tables(page.extract_tables() or [])
                    combined = f"{page_text}\n{table_text}".strip()
                    if combined:
                        all_pages[i] = combined
                except Exception as exc:
                    logger.warning("pdfplumber page %s failed: %s", i, exc)
    except Exception as exc:
        logger.warning("pdfplumber failed for %s: %s", filename, exc)

    # Strategy 2 — PyMuPDF in parallel for pages pdfplumber missed
    missing = _missing_pages(all_pages, total_pages)
    if missing:
        for page_num, text in _extract_pymupdf_pages_parallel(file_path, missing).items():
            all_pages[page_num] = text

    # Strategy 3 — PyPDF2 for remaining pages
    still_missing = _missing_pages(all_pages, total_pages)
    if still_missing:
        try:
            import PyPDF2

            with open(file_path, "rb") as handle:
                reader = PyPDF2.PdfReader(handle)
                for page_num in still_missing:
                    try:
                        text = (reader.pages[page_num - 1].extract_text() or "").strip()
                        if text:
                            all_pages[page_num] = text
                    except Exception as exc:
                        logger.warning("PyPDF2 page %s failed: %s", page_num, exc)
        except Exception as exc:
            logger.warning("PyPDF2 failed for %s: %s", filename, exc)

    # Strategy 4 — Gemini OCR for pages local extractors could not read
    ocr_candidates = _ocr_candidate_pages(all_pages, total_pages)
    ocr_limit = _ocr_page_limit(total_pages, len(ocr_candidates))
    if ocr_limit > 0:
        pages_to_ocr = ocr_candidates[:ocr_limit]
        if len(ocr_candidates) > ocr_limit:
            logger.warning(
                "%s pages need OCR in %s but capping at %s",
                len(ocr_candidates),
                filename,
                ocr_limit,
            )
        if on_progress:
            on_progress(f"OCR on {len(pages_to_ocr)} scanned page(s)…")

        for page_num, text in _gemini_ocr_pages_parallel(
            file_path, pages_to_ocr, on_progress=on_progress
        ).items():
            all_pages[page_num] = text

    _truncate_dense_pages(all_pages)

    output: List[str] = []
    for page_num in sorted(all_pages.keys()):
        content = all_pages[page_num].strip()
        if content:
            output.append(f"[PAGE {page_num}]\n{content}")

    unreadable = _missing_pages(all_pages, total_pages)
    if unreadable:
        logger.warning(
            "Unreadable pages after all strategies in %s: %s",
            filename,
            unreadable,
        )

    if not output:
        raise ValueError(
            f"No readable text found in '{filename}'. "
            "This PDF may be blank or image-only."
        )

    logger.info(
        "Deep PDF extraction for %s: %s/%s pages with text",
        filename,
        len(output),
        total_pages,
    )
    return "\n\n".join(output)


def _extract_non_pdf(file_path: str, filename: str) -> str:
    ext = Path(file_path).suffix.lower() or os.path.splitext(filename.lower())[1]

    if ext == ".docx":
        import docx

        doc = docx.Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        return f"[PAGE 1]\n" + "\n".join(paragraphs)

    with open(file_path, "r", encoding="utf-8", errors="ignore") as handle:
        body = handle.read()
    return f"[PAGE 1]\n{body}"


def extract_text_with_pages(
    file_path: str,
    filename: str,
    on_progress: Optional[ProgressCallback] = None,
) -> str:
    """
    Extract text preserving page numbers.
    PDFs use deep multi-strategy per-page extraction; other formats use [PAGE 1].
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found at path: {file_path}")

    row_chunks = parse_row_chunks(file_path, filename)
    if row_chunks:
        return f"[PAGE 1]\n" + "\n".join(row_chunks)

    ext = Path(file_path).suffix.lower() or os.path.splitext(filename.lower())[1]

    try:
        if ext == ".pdf":
            return _extract_pdf_deep(file_path, filename, on_progress=on_progress)
        return _extract_non_pdf(file_path, filename)
    except Exception as exc:
        raise ValueError(f"Error parsing file '{filename}': {exc}") from exc


def extract_text(
    file_path: str,
    filename: str,
    on_progress: Optional[ProgressCallback] = None,
) -> str:
    """Extract plain text with page markers for RAG indexing."""
    return extract_text_with_pages(file_path, filename, on_progress=on_progress)
